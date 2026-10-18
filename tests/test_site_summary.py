"""site_summary（敷地マネージャーのパレット相当の組み立て）のテスト。

データは実プロジェクト（BLCJ28RC庁舎サンプルモデル）で取得した値と、
敷地マネージャーのパレット表示（2352㎡ / 第二種住居地域 / 建蔽率 60% 等）に合わせている。
"""
import pytest
from openpyxl import load_workbook

import site_summary as ss
import table_tools

CATEGORIES = ["第一種低層住居専用地域", "第二種低層住居専用地域", "第一種中高層住居専用地域",
              "第二種中高層住居専用地域", "第一種住居地域", "第二種住居地域", "準住居地域", "田園住居地域",
              "近隣商業地域", "商業地域", "準工業地域", "工業地域", "工業専用地域", "指定なし"]

SITE = {
    "guid": "4E2A5AA0-2EB5-4534-8F80-C41356505D42", "libPart": "敷地オブジェクト", "floor": 0,
    "params": {
        "numVertex": 4, "vertexXArray": [-9, -9, 39, 39], "vertexYArray": [46, -3, -3, 46],
        "lengthBoundary": ["49.000", "48.000", "49.000", "48.000"],
        "iBoundaryTypeArray": [2, 1, 1, 1], "strBoundaryType": ["隣地境界線", "道路境界線", "その他"],
        "numAreaBoundary": 0, "iSiteCategory1": 6, "iSiteCategory2": 1,
        "strSiteCategory": CATEGORIES,
    },
    "descs": {"iBoundaryType": "道路境界線"},
}
ROAD = {
    "guid": "989863C2-59C0-4D6E-8C86-2D77626F9D00", "libPart": "道路境界線", "floor": 0,
    "params": {
        "numVertex": 2, "vertexXArray": [-41.7457, 77.3745], "vertexYArray": [52.082, 52.082],
        "strRoadName": "北側道路", "iRoadType": 1, "iRoadProperty": 1,
        "strRoadType": ["法42条1項1号", "法42条1項2号"], "strRoadProperty": ["道路中心線", "道路境界線", "指定なし"],
    },
    "descs": {"iBoundaryType": "道路境界線"},
}
LABEL = {"guid": "L1", "libPart": "道路ラベル", "params": {}, "descs": {}}
PROJECT = [
    {"desc": "敷地面積の合計", "key": "SITEAREA", "value": "2352.00"},
    {"desc": "建築面積建物物全体 (合計) ", "key": "BA", "value": "864.94㎡"},
    {"desc": "延べ面積 容積対象", "key": "FA", "value": "1,723.81"},
    {"desc": "許容建蔽率", "key": "BCR", "value": "60%"},
    {"desc": "許容容積率", "key": "FAR", "value": "160"},
    {"desc": "用途地域1", "key": "Z1", "value": "第二種住居地域"},
    {"desc": "プロジェクト名", "key": "PROJECTNAME", "value": "庁舎"},
]


def payload(**kw):
    p = {"type": "site_summary", "elements": [SITE, ROAD, LABEL], "projectInfo": PROJECT}
    p.update(kw)
    return p


# ── 小道具 ─────────────────────────────────────────────────────────

@pytest.mark.parametrize("v, expected", [
    ("864.94㎡", 864.94), ("1,723.81", 1723.81), ("-546.27", -546.27), (".5", 0.5),
    (12, 12.0), ("なし", None), (None, None), ("", None), (True, None),
])
def test_parse_number(v, expected):
    assert ss.parse_number(v) == expected


@pytest.mark.parametrize("v, expected", [("60%", 60.0), ("60", 60.0), ("0.6", 60.0), ("160", 160.0), ("", None)])
def test_parse_percent(v, expected):
    assert ss.parse_percent(v) == pytest.approx(expected) if expected is not None else ss.parse_percent(v) is None


def test_polygon_area():
    assert ss.polygon_area([-9, -9, 39, 39], [46, -3, -3, 46]) == pytest.approx(2352.0)
    assert ss.polygon_area([0, 1], [0, 1]) == 0.0


def test_formatters():
    assert ss.fmt_area(1411.2) == "1,411.20㎡" and ss.fmt_area(None) == "---"
    assert ss.fmt_percent(36.776) == "36.78%" and ss.fmt_len(48) == "48.000m"


# ── パレット相当の組み立て ─────────────────────────────────────────

def test_site_shape_boundaries_and_district():
    s = ss.build_site_summary(payload())
    assert len(s["sites"]) == 1 and len(s["roads"]) == 1     # ラベルは対象外
    site = s["sites"][0]
    assert site["area"] == pytest.approx(2352.0)
    assert site["perimeter"] == pytest.approx(194.0)
    assert [(b["no"], b["length"], b["type"]) for b in site["boundaries"]] == [
        (1, 49.0, "道路境界線"), (2, 48.0, "隣地境界線"), (3, 49.0, "隣地境界線"), (4, 48.0, "隣地境界線")]
    assert site["districts"] == [{"no": 1, "name": "第二種住居地域"}]    # iSiteCategory1=6
    assert s["districts"] == [["第二種住居地域", "2,352.00㎡"]]


def test_road():
    road = ss.build_site_summary(payload())["roads"][0]
    assert (road["name"], road["type"], road["property"], road["boundaryType"]) == \
        ("北側道路", "法42条1項1号", "道路中心線", "道路境界線")
    assert road["length"] == pytest.approx(119.1202)


def test_check_table_matches_palette():
    s = ss.build_site_summary(payload())
    assert s["totals"] == [["敷地面積の合計", "2,352.00㎡"], ["建築面積の合計", "864.94㎡"],
                           ["延床面積の合計", "1,723.81㎡"]]
    # プロジェクト情報に計画比率が無い場合は 面積÷敷地面積 で計算（864.94/2352 = 36.77%）
    assert s["check"] == [
        ["建蔽率", "60.00%", "36.77%", "適合"],
        ["建築面積", "1,411.20㎡", "864.94㎡", "-546.26㎡"],
        ["容積率", "160.00%", "73.29%", "適合"],
        ["容積対象面積", "3,763.20㎡", "1,723.81㎡", "-2,039.39㎡"],
    ]
    assert not s["missing_totals"]
    assert ["用途地域1", "第二種住居地域"] in s["project_rows"]
    assert ["プロジェクト名", "PROJECTNAME", "庁舎"] in s["project_all"]


def test_planned_ratio_prefers_project_info():
    """敷地マネージャーが書いた計画比率（丸め前の面積で計算）を優先し、パレット表示と一致させる。"""
    proj = PROJECT + [{"desc": "建蔽率", "value": "36.78%"}, {"desc": "容積率", "value": "73.30%"}]
    s = ss.build_site_summary(payload(projectInfo=proj))
    assert s["check"][0] == ["建蔽率", "60.00%", "36.78%", "適合"]
    assert s["check"][2] == ["容積率", "160.00%", "73.30%", "適合"]


def test_over_limit_is_not_compliant():
    proj = [dict(p) for p in PROJECT]
    proj[1]["value"] = "1500"        # 建築面積が許容 1411.20 を超える
    s = ss.build_site_summary(payload(projectInfo=proj))
    assert s["check"][0][3] == "不適合"
    assert s["check"][1][3] == "88.80㎡"


def test_fallbacks_without_project_info():
    """プロジェクト情報が空でも、敷地形状から面積だけは出す。集計値は未取得として扱う。"""
    s = ss.build_site_summary(payload(projectInfo=[]))
    assert s["totals"][0] == ["敷地面積の合計", "2,352.00㎡"]
    assert s["check"][0] == ["建蔽率", "---", "---", "---"]
    assert s["missing_totals"]


def test_specified_ratio_fallback_and_multi_district():
    site = {**SITE, "params": {**SITE["params"], "numAreaBoundary": 1, "iSiteCategory2": 10}}
    proj = [{"desc": "規定建蔽率1", "value": "0.6"}, {"desc": "規定容積率1", "value": "200%"},
            {"desc": "敷地面積1", "value": "1500"}, {"desc": "敷地面積2", "value": "852"}]
    s = ss.build_site_summary(payload(elements=[site], projectInfo=proj))
    assert s["districts"] == [["第二種住居地域", "1,500.00㎡"], ["商業地域", "852.00㎡"]]
    assert s["check"][0][1] == "60.00%" and s["check"][2][1] == "200.00%"


def test_malformed_params_are_tolerated():
    bad = {"guid": "X", "libPart": "敷地オブジェクト",
           "params": {"vertexXArray": [0, "a", 1], "vertexYArray": [0, 1], "iSiteCategory1": "x",
                      "iBoundaryTypeArray": [9], "strBoundaryType": []}}
    s = ss.build_site_summary({"elements": [bad], "projectInfo": []})
    assert s["sites"][0]["districts"] == [{"no": 1, "name": None}]
    assert s["districts"] == [["---", "---"]]
    assert ss.build_site_summary({})["sites"] == []


# ── Excel ──────────────────────────────────────────────────────────

def test_summary_sheets_export(tmp_path):
    path = tmp_path / "site.xlsx"
    table_tools.export_sheets(path, ss.summary_sheets(ss.build_site_summary(payload())))
    wb = load_workbook(path)
    assert wb.sheetnames == ["集計結果", "合計", "用途地域", "境界線", "道路", "プロジェクト情報"]
    assert [c.value for c in wb["集計結果"][2]] == ["建蔽率", "60.00%", "36.77%", "適合"]
    assert [c.value for c in wb["境界線"][2]] == [1, 1, 49, "道路境界線"]
    assert wb["道路"]["A2"].value == "北側道路"
