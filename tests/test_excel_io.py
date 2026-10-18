"""Excel 出力・取込を各ウィンドウで確実に行えることの確認。

- データ表示・編集ウィンドウ（条件設定から / 一覧表から開いたもの）
- 敷地情報（敷地マネージャー）: 出力と「プロジェクト情報」シートの取込（Archicad へ書き込み）
"""
import os

import pytest
from openpyxl import load_workbook
from PyQt6.QtWidgets import QMessageBox

import server
import table_tools

G = ["AAAAAAAA-1", "BBBBBBBB-2"]
PROPS = [
    {"guid": "P-W", "name": "幅", "group": "形状", "editable": True},
    {"guid": "P-F", "name": "耐火", "group": "一般", "editable": True, "enums": ["1時間", "2時間"]},
]


@pytest.fixture
def win(main_window, tmp_path, monkeypatch):
    w = main_window
    w.current_elements = [{"guid": g, "type": "Wall"} for g in G]
    w.current_props = list(PROPS)
    w.on_values({"values": [{"guid": g, "modiStamp": 1, "props": {"P-W": str(100 * (i + 1)), "P-F": "1時間"}}
                            for i, g in enumerate(G)]})
    w.data_panel.setWindowTitle("データ表示・編集 #1  1FL / Wall  (2件 × 2列)")
    return w


def _save_to(monkeypatch, path, seen=None):
    def fake(parent, title, start, flt):
        if seen is not None:
            seen.append(start)
        return str(path), ""
    monkeypatch.setattr(server.QFileDialog, "getSaveFileName", staticmethod(fake))


def _open(monkeypatch, path):
    monkeypatch.setattr(server.QFileDialog, "getOpenFileName", staticmethod(lambda *a, **k: (str(path), "")))


# ── table_tools ──────────────────────────────────────────────────────

def test_safe_file_name():
    assert table_tools.safe_file_name("一覧表: 内部仕上表 1階") == "一覧表_ 内部仕上表 1階.xlsx"
    assert table_tools.safe_file_name('a/b\\c*?"<>|') == "a_b_c.xlsx"   # 末尾の _ は取り除く
    assert table_tools.safe_file_name("") == "BIMデータ.xlsx"


def test_read_sheet_and_missing_sheet(tmp_path):
    p = tmp_path / "s.xlsx"
    table_tools.export_sheets(p, [("プロジェクト情報", ["項目", "キー", "値"], [["名前", "K1", "庁舎"], ["面積", "K2", "100"]])])
    assert table_tools.read_sheet(p, "プロジェクト情報") == [["項目", "キー", "値"], ["名前", "K1", "庁舎"], ["面積", "K2", "100"]]
    with pytest.raises(ValueError, match="見つかりません"):
        table_tools.read_sheet(p, "無いシート")


# ── データ表示・編集ウィンドウ ───────────────────────────────────────

def test_export_default_name_last_dir_and_verify(win, monkeypatch, tmp_path):
    seen = []
    _save_to(monkeypatch, tmp_path / "out", seen)            # 拡張子なし → .xlsx を付ける
    win.export_excel()
    assert os.path.basename(seen[0]) == "データ表示・編集 #1  1FL _ Wall.xlsx"   # ウィンドウ名から（件数部分は除く）
    assert (tmp_path / "out.xlsx").exists()
    assert table_tools.count_data_rows(tmp_path / "out.xlsx") == 2
    kind, _, text = win.msgbox.calls[-1]
    assert kind == "information" and "2 行 × 2 列" in text
    seen.clear()
    win.export_excel()
    assert os.path.dirname(seen[0]) == str(tmp_path)          # 前回のフォルダを記憶


def test_export_reports_hidden_rows_and_mismatch(win, monkeypatch, tmp_path):
    _save_to(monkeypatch, tmp_path / "f.xlsx")
    win.search_edit.setText("AAAA")
    win.export_excel()
    assert "非表示の 1 行" in win.msgbox.calls[-1][2]
    monkeypatch.setattr(table_tools, "count_data_rows", lambda p: 99)
    win.export_excel()
    assert win.msgbox.calls[-1][0] == "warning" and "一致しません" in win.msgbox.calls[-1][2]


def test_import_confirm_decline_and_no_changes(win, monkeypatch, tmp_path):
    path = tmp_path / "e.xlsx"
    _save_to(monkeypatch, path)
    win.export_excel()
    _open(monkeypatch, path)
    win.import_excel()                                        # 何も編集していない
    assert "変更）はありませんでした" in win.msgbox.calls[-1][2]

    wb = load_workbook(path); ws = wb[table_tools.DATA_SHEET]
    ws["B2"] = 150; ws["C3"] = "2時間"; wb.save(path)
    win.msgbox.question_answer = QMessageBox.StandardButton.No
    win.import_excel()                                        # 確認で「いいえ」→ 反映しない
    assert win.model._values[(G[0], "P-W")] == "100"
    assert win.status_label.text() == "Excel取込: 中止しました"
    confirm = [c for c in win.msgbox.calls if c[0] == "question"][-1][2]
    assert "2 セルを表に反映" in confirm and "幅: 1 セル" in confirm and "耐火: 1 セル" in confirm

    win.msgbox.question_answer = QMessageBox.StandardButton.Yes
    win.import_excel()
    assert win.model._values[(G[0], "P-W")] == "150" and win.model._values[(G[1], "P-F")] == "2時間"
    assert win.undo_stack.count() == 1


def test_import_locked_file_message(win, monkeypatch, tmp_path):
    _open(monkeypatch, tmp_path / "x.xlsx")
    monkeypatch.setattr(table_tools, "read_xlsx", lambda *a: (_ for _ in ()).throw(PermissionError("locked")))
    win.import_excel()
    assert "閉じてから" in win.msgbox.calls[-1][2]


def test_import_cancel_dialog(win, monkeypatch):
    monkeypatch.setattr(server.QFileDialog, "getOpenFileName", staticmethod(lambda *a, **k: ("", "")))
    win.import_excel()
    assert win.msgbox.calls == []


def test_schedule_window_excel_roundtrip(main_window, monkeypatch, tmp_path):
    """一覧表から開いたウィンドウも、見出し（建具主寸法 / W 等）のまま出力・取込できる。"""
    from tests.test_schedules import _open_door_schedule, folder as _folder_fixture  # noqa: F401
    import tests.test_schedules as TS
    d = tmp_path / "連動一覧表"; d.mkdir()
    TS.scheme(str(d), 18, "建具表_ドア", [TS.crit("88", 1, cls=TS.DOOR)],
              [TS.field(1, 0, "表", ptype="0"), TS.field(2, 1, "", name="建具種類", group="建具情報(共通)", desc="建具種類"),
               TS.field(4, 1, "建具主寸法", ptype="0"), TS.field(5, 4, "W", guid=TS.P_WIDTH)])
    monkeypatch.setenv("BIM_SCHEDULE_DIR", str(d))
    w = main_window
    dlg = w.show_schedule_list(); dlg.tbl.selectRow(0); dlg.btn_open.click()
    w.on_schedule_resolved(TS._resolve_reply(w.sent.last("resolve_schedule")))
    w.on_elements({"reqId": w.sent.last("get_elements")["reqId"], "elements": [{"guid": "D1", "type": "Door"}]})
    w.on_values({"reqId": w.sent.last("get_values")["reqId"], "values": [
        {"guid": "D1", "props": {"P-建具種類": "AD", TS.P_WIDTH: "0.9"}}]})
    path = tmp_path / "sched.xlsx"
    seen = []
    _save_to(monkeypatch, path, seen)
    w.export_excel()
    assert os.path.basename(seen[0]).startswith("一覧表_ 建具表_ドア")
    ws = load_workbook(path)[table_tools.DATA_SHEET]
    assert [c.value for c in ws[1]] == ["GUID", "建具種類", "建具主寸法 / W"]
    wb = load_workbook(path); wb[table_tools.DATA_SHEET]["B2"] = "SS"; wb.save(path)
    _open(monkeypatch, path)
    w.import_excel()
    assert w.model._values[("D1", "P-建具種類")] == "SS"


# ── 敷地情報（敷地マネージャー）─────────────────────────────────────

def _site(w):
    from tests.test_site_summary import payload
    w.on_site_summary(payload())
    return w._site_dlg


def test_site_export_verify_and_import_project_info(main_window, monkeypatch, tmp_path):
    w = main_window
    dlg = _site(w)
    path = tmp_path / "site.xlsx"
    seen = []
    _save_to(monkeypatch, path, seen)
    dlg.btn_export.click()
    assert os.path.basename(seen[0]) == "敷地情報.xlsx"
    assert "シートを書き出しました" in w.msgbox.calls[-1][2]

    # 「プロジェクト情報」シートで 許容建蔽率 と プロジェクト名 を編集
    wb = load_workbook(path); ws = wb["プロジェクト情報"]
    for row in ws.iter_rows(min_row=2):
        if row[1].value == "BCR":
            row[2].value = "50%"
        if row[1].value == "PROJECTNAME":
            row[2].value = "新庁舎"
    ws.append(["謎", "NOKEY", "x"])
    wb.save(path)
    _open(monkeypatch, path)
    dlg.btn_import.click()
    confirm = w.msgbox.calls[-1]
    assert confirm[0] == "question" and "2 件更新" in confirm[2]
    assert "許容建蔽率: 60% → 50%" in confirm[2] and "プロジェクト名: 庁舎 → 新庁舎" in confirm[2]
    assert "プロジェクト情報に無い項目（無視）: 謎" in confirm[2]
    cmd = w.sent.last("set_project_info")
    keys = [bytes.fromhex(h).decode() for h in cmd["keysHex"]]
    vals = [bytes.fromhex(h).decode() for h in cmd["valuesHex"]]
    assert dict(zip(keys, vals)) == {"BCR": "50%", "PROJECTNAME": "新庁舎"}

    w.on_project_info_result({"type": "project_info_result", "set": 2, "failed": []})
    assert "2 件を Archicad に書き込みました" in w.msgbox.calls[-1][2]
    assert w.sent.commands()[-1] == "get_site_summary"           # 最新値で表示し直す
    w.on_project_info_result({"set": 1, "failed": ["PROJECTNAME"]})
    assert w.msgbox.calls[-1][0] == "warning" and "PROJECTNAME" in w.msgbox.calls[-1][2]


def test_site_import_no_change_decline_and_errors(main_window, monkeypatch, tmp_path):
    w = main_window
    dlg = _site(w)
    path = tmp_path / "site.xlsx"
    _save_to(monkeypatch, path)
    dlg.export_excel()
    _open(monkeypatch, path)
    dlg.import_excel()
    assert "変更はありませんでした" in w.msgbox.calls[-1][2]

    wb = load_workbook(path); ws = wb["プロジェクト情報"]; ws["C2"] = "9999"; wb.save(path)
    w.msgbox.question_answer = QMessageBox.StandardButton.No
    dlg.import_excel()
    assert "set_project_info" not in w.sent.commands() and "中止" in dlg.info_label.text()

    other = tmp_path / "other.xlsx"
    table_tools.export_sheets(other, [("別", ["a"], [])])
    _open(monkeypatch, other)
    dlg.import_excel()
    assert w.msgbox.calls[-1][0] == "warning" and "プロジェクト情報" in w.msgbox.calls[-1][2]

    _open(monkeypatch, path)
    monkeypatch.setattr(table_tools, "read_sheet", lambda *a: (_ for _ in ()).throw(PermissionError("x")))
    dlg.import_excel()
    assert "閉じてから" in w.msgbox.calls[-1][2]

    monkeypatch.setattr(server.QFileDialog, "getOpenFileName", staticmethod(lambda *a, **k: ("", "")))
    n = len(w.msgbox.calls)
    dlg.import_excel()
    assert len(w.msgbox.calls) == n
    dlg._summary = None
    dlg.import_excel(); dlg.export_excel()                     # 未取得なら何もしない


def test_site_export_locked_and_unverified(main_window, monkeypatch, tmp_path):
    w = main_window
    dlg = _site(w)
    _save_to(monkeypatch, tmp_path / "s.xlsx")
    monkeypatch.setattr(table_tools, "read_sheet", lambda *a: [["h"]])
    dlg.export_excel()
    assert "確認できませんでした" in w.msgbox.calls[-1][2]
    monkeypatch.setattr(table_tools, "export_sheets", lambda *a: (_ for _ in ()).throw(PermissionError("x")))
    dlg.export_excel()
    assert "閉じてから" in w.msgbox.calls[-1][2]


# ── Excel 出力設定: GUID 列を編集不可 / 非表示 ──────────────────────

def test_export_locks_and_hides_guid_by_default(win, monkeypatch, tmp_path):
    path = tmp_path / "g.xlsx"
    _save_to(monkeypatch, path)
    win.export_excel()
    ws = load_workbook(path)[table_tools.DATA_SHEET]
    assert ws.column_dimensions["A"].hidden is True                  # GUID 列は非表示
    assert ws.protection.sheet is True                               # シート保護
    assert ws["A2"].protection.locked is True and ws["A1"].protection.locked is True   # GUID・見出しは編集不可
    assert ws["B2"].protection.locked is False and ws["C3"].protection.locked is False  # 値は編集できる
    assert ws.protection.autoFilter is False                         # 保護中もフィルターは使える
    assert ws.protection.password is None                            # パスワードなし（必要なら解除できる）
    text = win.msgbox.calls[-1][2]
    assert "GUID 列と見出し行は編集不可" in text and "GUID 列は非表示" in text

    # 非表示・保護のまま値を編集して取り込める
    wb = load_workbook(path); wb[table_tools.DATA_SHEET]["B2"] = 175; wb.save(path)
    _open(monkeypatch, path)
    win.import_excel()
    assert win.model._values[(G[0], "P-W")] == "175"


def test_export_options_can_be_turned_off_and_are_remembered(win, monkeypatch, tmp_path):
    p = win.data_panel
    assert p.act_lock_guid.isChecked() and p.act_hide_guid.isChecked()   # 既定は ON
    p.act_lock_guid.setChecked(False)
    p.act_hide_guid.setChecked(False)
    assert win.excel_option(server.EXCEL_LOCK_GUID_KEY) is False
    path = tmp_path / "plain.xlsx"
    _save_to(monkeypatch, path)
    win.export_excel()
    ws = load_workbook(path)[table_tools.DATA_SHEET]
    assert ws.column_dimensions["A"].hidden is False and ws.protection.sheet is False
    text = win.msgbox.calls[-1][2]
    assert "編集不可（シート保護" not in text and "GUID 列は非表示" not in text   # 設定 OFF なら注記も出ない

    other = win._create_data_panel()                                 # 別のウィンドウも同じ設定
    assert not other.act_lock_guid.isChecked() and not other.act_hide_guid.isChecked()
    p.act_hide_guid.setChecked(True)
    other.btn_excel_opts.menu().aboutToShow.emit()                   # メニューを開くと最新の設定を反映
    assert other.act_hide_guid.isChecked()


def test_table_tools_export_options(tmp_path):
    props = [{"guid": "P", "name": "幅"}]
    path = tmp_path / "t.xlsx"
    table_tools.export_xlsx(path, props, [("G1", ["1"])], lock_guid=True, hide_guid=False)
    ws = load_workbook(path)[table_tools.DATA_SHEET]
    assert ws.protection.sheet and not ws.column_dimensions["A"].hidden
    table_tools.export_xlsx(path, props, [("G1", ["1"])])            # 既定（関数としては OFF）
    ws = load_workbook(path)[table_tools.DATA_SHEET]
    assert not ws.protection.sheet and ws["B2"].protection.locked is True