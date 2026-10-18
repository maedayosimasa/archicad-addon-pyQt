"""データ表示・編集テーブルの Excel 風機能（並び替え・検索・コピー/貼り付け・一括入力・Excel 入出力）。"""
import pytest
from openpyxl import load_workbook
from PyQt6.QtCore import QItemSelection, QItemSelectionModel, QPoint, Qt
from PyQt6.QtWidgets import QApplication

import server
import table_tools

G = ["AAAAAAAA-1", "BBBBBBBB-2", "CCCCCCCC-3", "DDDDDDDD-4"]
PROPS = [
    {"guid": "builtin:element_id", "name": "ID", "group": "IDとカテゴリ", "editable": True},
    {"guid": "P-W", "name": "幅", "group": "形状", "editable": True},
    {"guid": "P-F", "name": "耐火", "group": "一般", "editable": True, "enums": ["1時間", "2時間"]},
    {"guid": "P-A", "name": "面積", "group": "一般", "editable": False},
]
WIDTHS = ["300", "20", "", "1000"]


@pytest.fixture
def win(main_window):
    w = main_window
    w.current_elements = [{"guid": g, "type": "Wall"} for g in G]
    w.current_props = list(PROPS)
    w.on_values({"values": [
        {"guid": g, "modiStamp": i + 1,
         "props": {"builtin:element_id": f"W-{i}", "P-W": WIDTHS[i], "P-F": "1時間", "P-A": "5"}}
        for i, g in enumerate(G)
    ]})
    w.sent.clear()
    return w


def col_values(w, col):
    return [w.proxy.index(r, col).data() for r in range(w.proxy.rowCount())]


def select(w, cells):
    """(表示行, 列) の一覧を選択状態にする。"""
    sm = w.table.selectionModel()
    # setCurrentIndex は既定で選択をクリアするため、NoUpdate で先にカーソルだけ移動する
    sm.setCurrentIndex(w.proxy.index(*cells[0]), QItemSelectionModel.SelectionFlag.NoUpdate)
    sm.clearSelection()
    for r, c in cells:
        idx = w.proxy.index(r, c)
        sm.select(QItemSelection(idx, idx), QItemSelectionModel.SelectionFlag.Select)


# ── 並び替え・検索 ─────────────────────────────────────────────────

def test_initial_order_is_fetch_order_and_row_count(win):
    assert col_values(win, 2) == WIDTHS
    assert win.table.isSortingEnabled()
    assert win.row_count_label.text() == "表示 4 / 全 4 件"


def test_header_sort_numeric_and_blanks_last(win):
    win.table.sortByColumn(2, Qt.SortOrder.AscendingOrder)
    assert col_values(win, 2) == ["20", "300", "1000", ""]      # 文字列順なら "1000" < "20"
    win.table.sortByColumn(2, Qt.SortOrder.DescendingOrder)
    assert col_values(win, 2) == ["1000", "300", "20", ""]      # 降順でも空欄は末尾
    win.clear_sort()
    assert col_values(win, 2) == WIDTHS
    assert win.table.horizontalHeader().sortIndicatorSection() == -1


def test_search_filters_all_columns(win):
    win.search_edit.setText("w-2")
    assert col_values(win, 1) == ["W-2"]
    assert win.row_count_label.text() == "表示 1 / 全 4 件"
    win.search_edit.setText("CCCC")                       # ID 列（GUID 先頭）も対象
    assert col_values(win, 0) == ["CCCCCCCC"]
    win.search_edit.clear()
    assert win.proxy.rowCount() == 4


def test_click_on_sorted_row_selects_right_element(win):
    win.table.sortByColumn(2, Qt.SortOrder.AscendingOrder)   # 先頭行は幅 20 の G[1]
    win.on_table_clicked(win.proxy.index(0, 2))
    assert win.sent.last("select_elements")["guids"] == [G[1]]


# ── コピー ─────────────────────────────────────────────────────────

def test_copy_selection_as_tsv(win, qapp):
    select(win, [(0, 1), (0, 2), (1, 1), (1, 2)])
    win.copy_selection()
    assert QApplication.clipboard().text() == "W-0\t300\nW-1\t20\n"
    assert win.status_label.text() == "コピー: 4 セル"


def test_copy_sparse_selection_fills_gaps(win, qapp):
    select(win, [(0, 1), (2, 2)])
    win.copy_selection()
    assert table_tools.parse_tsv(QApplication.clipboard().text()) == [["W-0", ""], ["", ""]]


def test_copy_without_selection_does_nothing(win, qapp):
    QApplication.clipboard().setText("keep")
    win.table.selectionModel().clearSelection()
    win.table.setCurrentIndex(win.proxy.index(-1, -1))
    win.copy_selection()
    assert QApplication.clipboard().text() == "keep"


# ── 貼り付け（検証・Undo 1 回で戻る）──────────────────────────────

def test_paste_block_from_excel_and_undo_as_one_step(win, qapp):
    QApplication.clipboard().setText("111\t2時間\r\n222\t1時間\r\n")
    select(win, [(1, 2)])
    win.paste_clipboard()
    v = win.model._values
    assert (v[(G[1], "P-W")], v[(G[1], "P-F")], v[(G[2], "P-W")]) == ("111", "2時間", "222")
    assert v[(G[2], "P-F")] == "1時間"                                   # 同じ値は変更しない
    assert win.model.index(1, 2).data(Qt.ItemDataRole.ForegroundRole).name() == "#ff0000"
    assert win.undo_stack.count() == 1 and win.undo_stack.text(0).startswith("貼り付け (3 セル)")
    assert len(win.db.get_history()) == 3

    win._do_undo()
    assert (v[(G[1], "P-W")], v[(G[1], "P-F")], v[(G[2], "P-W")]) == ("20", "1時間", "")


def test_paste_single_value_fills_selection(win, qapp):
    QApplication.clipboard().setText("2時間")
    select(win, [(0, 3), (1, 3), (3, 3)])
    win.paste_clipboard()
    assert [win.model._values[(g, "P-F")] for g in G] == ["2時間", "2時間", "1時間", "2時間"]


def test_paste_validation_skips_and_warns(win, qapp):
    # 列: 0=ID表示(編集不可) 1=IDプロパティ(同値) 2=幅(同値) 3=耐火(列挙外) 4=面積(編集不可) 5=表の外
    QApplication.clipboard().setText("X\tW-0\t300\t3時間\t99\tはみ出し")
    select(win, [(0, 0)])
    win.paste_clipboard()
    kind, title, text = win.msgbox.calls[-1]
    assert kind == "warning" and "スキップ" in title
    assert "ID 列は編集できません" in text and "「3時間」は候補にありません" in text and "編集不可" in text
    assert win.undo_stack.count() == 0                     # 何も変わらなければ Undo 履歴も増えない
    assert win.status_label.text() == "貼り付け: 0 セル変更 / 3 セルはスキップ"


def test_paste_into_sorted_view_hits_visible_rows(win, qapp):
    win.table.sortByColumn(2, Qt.SortOrder.AscendingOrder)   # 表示: G1(20), G0(300), G3(1000), G2("")
    QApplication.clipboard().setText("A\nB\n")
    select(win, [(0, 2)])
    win.paste_clipboard()
    assert win.model._values[(G[1], "P-W")] == "A"
    assert win.model._values[(G[0], "P-W")] == "B"


def test_paste_empty_clipboard_noop(win, qapp):
    QApplication.clipboard().setText("")
    select(win, [(0, 2)])
    win.paste_clipboard()
    assert win.undo_stack.count() == 0


def test_many_skip_reasons_are_summarized(win, qapp):
    targets = [(0, 3, f"候補外{i}") for i in range(12)]
    changed, skipped = win._apply_edits(targets, "テスト")
    assert changed == 0 and len(skipped) == 12
    assert "… 他 2 種類" in win.msgbox.calls[-1][2]


# ── 下方向へコピー・一括入力 ──────────────────────────────────────

def test_fill_down(win):
    select(win, [(0, 2), (1, 2), (2, 2), (0, 3), (1, 3)])
    win.fill_down()
    assert [win.model._values[(g, "P-W")] for g in G[:3]] == ["300", "300", "300"]
    assert win.undo_stack.count() == 1


def test_fill_down_needs_two_cells(win):
    select(win, [(0, 2)])
    win.fill_down()
    assert "2 セル以上" in win.status_label.text()
    assert win.undo_stack.count() == 0


def test_bulk_input_text(win, monkeypatch):
    monkeypatch.setattr(server.QInputDialog, "getText", staticmethod(lambda *a, **k: ("999", True)))
    select(win, [(0, 2), (2, 2), (0, 0)])                  # ID 列は対象外
    win.bulk_input()
    assert win.model._values[(G[0], "P-W")] == "999" and win.model._values[(G[2], "P-W")] == "999"


def test_bulk_input_enum_offers_choices(win, monkeypatch):
    seen = {}

    def get_item(parent, title, label, items, current, editable):
        seen["items"] = items
        return "2時間", True

    monkeypatch.setattr(server.QInputDialog, "getItem", staticmethod(get_item))
    select(win, [(0, 3), (1, 3)])
    win.bulk_input()
    assert seen["items"] == ["1時間", "2時間"]
    assert win.model._values[(G[1], "P-F")] == "2時間"


def test_bulk_input_cancel_and_empty(win, monkeypatch):
    monkeypatch.setattr(server.QInputDialog, "getText", staticmethod(lambda *a, **k: ("x", False)))
    select(win, [(0, 2)])
    win.bulk_input()
    select(win, [(0, 0)])
    win.bulk_input()                                        # ID 列のみ → 何もしない
    assert win.undo_stack.count() == 0


def test_cell_context_menu_runs_actions(win, monkeypatch, qapp):
    picked = []

    def fake_exec(menu, *a):
        texts = [a.text() for a in menu.actions()]
        picked.append(texts)
        next(a for a in menu.actions() if a.text().startswith("コピー")).trigger()

    monkeypatch.setattr(server.QMenu, "exec", fake_exec)
    select(win, [(0, 2)])
    win.show_cell_menu(QPoint(5, 5))
    assert picked[0][:4] == ["コピー\tCtrl+C", "貼り付け\tCtrl+V", "下方向へコピー\tCtrl+D", "選択セルに一括入力..."]
    assert QApplication.clipboard().text() == "300\n"


def test_shortcuts_registered_on_table(win):
    from PyQt6.QtGui import QShortcut
    keys = {s.key().toString() for s in win.table.findChildren(QShortcut)}
    assert {"Ctrl+C", "Ctrl+V", "Ctrl+D"} <= keys


# ── Excel 入出力 ─────────────────────────────────────────────────

def test_export_excel_visible_rows_with_red_changes(win, monkeypatch, tmp_path):
    path = tmp_path / "out.xlsx"
    monkeypatch.setattr(server.QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(path), "")))
    win.model.setData(win.model.index(0, 2), "350")
    win.search_edit.setText("W-0")
    win.export_excel()
    ws = load_workbook(path)[table_tools.DATA_SHEET]
    assert [c.value for c in ws[1]] == ["GUID", "ID", "幅", "耐火", "面積"]
    assert ws.max_row == 2 and ws["A2"].value == G[0] and ws["C2"].value == 350
    assert ws["C2"].fill.start_color.rgb.endswith("FFC7CE")
    assert "Excel出力: 1 行" in win.status_label.text()


def test_export_excel_cancel_empty_and_error(main_window, win, monkeypatch, tmp_path):
    monkeypatch.setattr(server.QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: ("", "")))
    win.export_excel()                                    # キャンセル
    monkeypatch.setattr(server.QFileDialog, "getSaveFileName",
                        staticmethod(lambda *a, **k: (str(tmp_path / "x.xlsx"), "")))
    monkeypatch.setattr(table_tools, "export_xlsx", lambda *a, **k: (_ for _ in ()).throw(PermissionError("開いています")))
    win.export_excel()
    assert win.msgbox.calls[-1][0] == "warning" and "閉じてから" in win.msgbox.calls[-1][2]   # Excel で開いている


def test_export_import_without_data(main_window):
    main_window.export_excel()
    main_window.import_excel()
    assert main_window.msgbox.kinds() == ["information", "information"]


def test_excel_roundtrip_import_applies_only_diffs(win, monkeypatch, tmp_path):
    path = tmp_path / "edit.xlsx"
    monkeypatch.setattr(server.QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(path), "")))
    win.export_excel()

    # Excel で編集したつもり: 幅を変更、耐火を変更、編集不可の面積を変更、無関係な行を追加
    wb = load_workbook(path)
    ws = wb[table_tools.DATA_SHEET]
    ws["C2"] = 301          # G0 幅 300 → 301
    ws["C3"] = 20.0         # G1 幅 20 → 20.0（数値として同じなので変更なし）
    ws["D4"] = "2時間"      # G2 耐火
    ws["E5"] = 77           # G3 面積（編集不可 → スキップ）
    ws.append(["ZZZZ-NOT-IN-TABLE", "x", 1, "1時間", 1])
    wb.save(path)

    monkeypatch.setattr(server.QFileDialog, "getOpenFileName", staticmethod(lambda *a, **k: (str(path), "")))
    win.import_excel()
    v = win.model._values
    assert v[(G[0], "P-W")] == "301" and v[(G[1], "P-W")] == "20" and v[(G[2], "P-F")] == "2時間"
    assert v[(G[3], "P-A")] == "5"
    assert win.undo_stack.count() == 1 and win.undo_stack.text(0).startswith("Excel取込 (2 セル)")
    confirm = next(c for c in win.msgbox.calls if c[0] == "question")
    assert "2 セルを表に反映" in confirm[2] and "取り込めないセル: 1" in confirm[2]   # 面積（編集不可）は確認時に表示
    assert win.msgbox.calls[-1][0] == "information"

    # 取込後の [Archicadへ反映] は変更セルのみを送る
    win.sync_data()
    req = win.sent.last("get_values")
    assert set(req["guids"]) == {G[0], G[2]}


def test_import_reports_unknown_columns_and_missing_rows(win, monkeypatch, tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["GUID", "幅", "謎の列"])
    wb.active.append([G[0], 555, "?"])
    wb.active.append(["NOPE", 1, "?"])
    path = tmp_path / "hand.xlsx"
    wb.save(path)
    monkeypatch.setattr(server.QFileDialog, "getOpenFileName", staticmethod(lambda *a, **k: (str(path), "")))
    win.import_excel()
    confirm = next(c for c in win.msgbox.calls if c[0] == "question")
    assert "1 セルを表に反映" in confirm[2]
    assert "表に無い要素: 1 行" in confirm[2] and "謎の列" in confirm[2]   # 注意点は反映前の確認に表示
    kind, _, text = win.msgbox.calls[-1]
    assert kind == "information" and "1 セルを変更" in text


def test_import_cancel_and_broken_file(win, monkeypatch, tmp_path):
    monkeypatch.setattr(server.QFileDialog, "getOpenFileName", staticmethod(lambda *a, **k: ("", "")))
    win.import_excel()
    assert win.msgbox.calls == []
    bad = tmp_path / "bad.xlsx"
    bad.write_text("not an excel file")
    monkeypatch.setattr(server.QFileDialog, "getOpenFileName", staticmethod(lambda *a, **k: (str(bad), "")))
    win.import_excel()
    assert win.msgbox.calls[-1][0] == "warning" and "読み込みに失敗" in win.msgbox.calls[-1][2]


# ── 検索と置換 ─────────────────────────────────────────────────────

def opts(find, replace="", **kw):
    o = {"find": find, "replace": replace, "match_case": False, "whole_cell": False,
         "regex": False, "selection_only": False}
    o.update(kw)
    return o


def test_replace_all_one_undo_step(win):
    changed, skipped, locked = win.replace_all(opts("W-", "壁-"))
    assert changed == 4 and skipped == 0 and locked == 0
    assert [win.model._values[(g, "builtin:element_id")] for g in G] == ["壁-0", "壁-1", "壁-2", "壁-3"]
    assert win.undo_stack.count() == 1 and win.undo_stack.text(0).startswith("置換 (4 セル)")
    win._do_undo()
    assert win.model._values[(G[0], "builtin:element_id")] == "W-0"


def test_replace_skips_id_column_and_locked_props(win):
    # "5" は 面積(編集不可) と ID 列以外にも… 幅には無い。ID 表示列 "CCCCCCCC" は対象外
    changed, skipped, locked = win.replace_all(opts("5", "6"))
    assert changed == 0 and locked == 4                  # 面積 4 セルは編集不可
    assert "編集不可で対象外 4 セル" in win.status_label.text()
    changed, _, locked = win.replace_all(opts("CCCC", "X"))
    assert changed == 0 and locked == 1                  # ID 表示列


def test_replace_enum_validation(win):
    changed, skipped, _ = win.replace_all(opts("1時間", "3時間"))  # 候補外 → 全スキップ
    assert changed == 0 and skipped == 4
    assert win.msgbox.calls[-1][0] == "warning"
    changed, skipped, _ = win.replace_all(opts("1", "2"))         # "1時間"→"2時間" は OK, 幅 1000→2000
    assert win.model._values[(G[0], "P-F")] == "2時間"
    assert win.model._values[(G[3], "P-W")] == "2000"


def test_replace_whole_cell_blank_and_regex(win):
    win.replace_all(opts("", "0", whole_cell=True))
    assert win.model._values[(G[2], "P-W")] == "0"               # 空欄だった幅
    win.replace_all(opts(r"^W-(\d)$", r"WALL\1", regex=True))
    assert win.model._values[(G[1], "builtin:element_id")] == "WALL1"


def test_replace_selection_only(win):
    select(win, [(0, 1), (1, 1)])
    changed, _, _ = win.replace_all(opts("W-", "X-", selection_only=True))
    assert changed == 2
    assert win.model._values[(G[2], "builtin:element_id")] == "W-2"


def test_replace_no_match_and_invalid_regex(win):
    assert win.replace_all(opts("存在しない", "x")) == (0, 0, 0)
    assert win.status_label.text() == "置換: 一致するセルはありません"
    assert win.replace_all(opts("(", "x", regex=True)) is None
    assert win.select_matches(opts("(", "x", regex=True)) is None
    assert win.msgbox.kinds()[-1] == "warning"


def test_select_matches(win):
    assert win.select_matches(opts("時間")) == 4
    sel = {(i.row(), i.column()) for i in win.table.selectionModel().selectedIndexes()}
    assert sel == {(r, 3) for r in range(4)}
    assert win.status_label.text() == "検索: 4 セル一致"
    assert win.select_matches(opts("なし")) == 0


def test_replace_dialog_flow(win):
    win.search_edit.setText("W-1")
    dlg = win.show_replace_dialog()
    assert dlg.isVisible() and dlg.find_edit.text() == "W-1"     # 検索欄の文字を引き継ぐ
    assert win.show_replace_dialog() is dlg                       # 2 回目は同じダイアログ
    dlg.find_edit.setText("W-1"); dlg.repl_edit.setText("壁1")
    dlg.chk_whole.setChecked(True)
    dlg.btn_find.click()
    assert dlg.info_label.text() == "一致: 1 セル"
    dlg.btn_replace.click()
    assert dlg.info_label.text() == "置換: 1 セル"
    assert win.model._values[(G[1], "builtin:element_id")] == "壁1"
    win.search_edit.clear()
    dlg.find_edit.setText("5"); dlg.repl_edit.setText("6"); dlg.chk_whole.setChecked(False)
    dlg.btn_replace.click()
    assert dlg.info_label.text() == "置換: 0 セル / 編集不可で対象外: 4"
    win.replace_all(opts("1時間", "3時間"))
    dlg.find_edit.setText("1時間"); dlg.repl_edit.setText("3時間")
    dlg.btn_replace.click()
    assert "検証でスキップ: 4" in dlg.info_label.text()
    dlg.chk_regex.setChecked(True); dlg.find_edit.setText("(")
    dlg.btn_find.click(); dlg.btn_replace.click()                 # 不正な正規表現 → 表示は更新しない
    assert "検証でスキップ: 4" in dlg.info_label.text()
    dlg.btn_close.click()
    assert not dlg.isVisible()


def test_replace_menu_and_shortcut(win, monkeypatch):
    from PyQt6.QtGui import QShortcut
    assert "Ctrl+H" in {s.key().toString() for s in win.table.findChildren(QShortcut)}
    texts = []
    monkeypatch.setattr(server.QMenu, "exec", lambda menu, *a: texts.extend(a.text() for a in menu.actions()))
    win.show_cell_menu(QPoint(1, 1))
    assert "検索と置換...\tCtrl+H" in texts


def test_excel_like_table_style(win):
    """タイトル行は少し濃い色、データ行は 1 行おきに薄い色（縞模様）。"""
    t = win.table
    assert t.alternatingRowColors()
    css = t.styleSheet()
    assert server.TABLE_HEADER_COLOR in css and "QHeaderView::section:horizontal" in css
    assert f"alternate-background-color: {server.TABLE_ALTERNATE_COLOR}" in css
    # 色付きセル（変更・確認済など）はモデルの色が優先され、縞模様で上書きされない
    win.model._change_status[G[0]] = 2
    assert win.model.index(0, 1).data(Qt.ItemDataRole.BackgroundRole).name() == "#c8f0c8"
    assert win.model.index(1, 1).data(Qt.ItemDataRole.BackgroundRole) is None   # 無色 → 縞模様が出る
    # 新しく開いたデータウィンドウにも同じ見た目
    other = win._create_data_panel()
    assert other.table.alternatingRowColors() and other.table.styleSheet() == css