"""table_tools（TSV・検証・並び替えキー・Excel 入出力）の単体テスト。"""
import pytest
from openpyxl import Workbook, load_workbook

import table_tools as tt

PROPS = [
    {"guid": "P-W", "name": "幅"},
    {"guid": "P-F", "name": "耐火", "enums": ["1時間", "2時間"]},
    {"guid": "P-RO", "name": "面積", "editable": False},
]


# ── TSV ─────────────────────────────────────────────────────────────

def test_tsv_roundtrip_with_special_cells():
    rows = [["a", "b\tc", None], ["改行\nあり", '引用"符', "3"]]
    text = tt.to_tsv(rows)
    assert tt.parse_tsv(text) == [["a", "b\tc", ""], ["改行\nあり", '引用"符', "3"]]


@pytest.mark.parametrize("text, expected", [
    ("", []),
    ("100", [["100"]]),
    ("100\r\n", [["100"]]),                       # Excel は末尾に改行を付ける
    ("1\t2\r\n3\t4\r\n", [["1", "2"], ["3", "4"]]),
    ("a\n\nb", [["a"], [], ["b"]]),
])
def test_parse_tsv(text, expected):
    assert tt.parse_tsv(text) == expected


# ── 検証 ───────────────────────────────────────────────────────────

def test_validate_value():
    assert tt.validate_value(PROPS[0], "123") == (True, "")
    assert tt.validate_value(PROPS[1], "2時間") == (True, "")
    ok, reason = tt.validate_value(PROPS[1], "3時間")
    assert not ok and "候補にありません" in reason
    ok, reason = tt.validate_value(PROPS[2], "1")
    assert not ok and "編集不可" in reason


def test_is_editable_default_true():
    assert tt.is_editable({}) and not tt.is_editable({"editable": False})


@pytest.mark.parametrize("a, b, eq", [
    ("100", "100", True), ("100.50", "100.5", True), ("100", 100, True),
    ("abc", "abd", False), ("1", "", False), ("", "", True),
])
def test_values_equal(a, b, eq):
    assert tt.values_equal(a, b) is eq


def test_sort_key_orders_numbers_text_blanks():
    vals = ["b", "10", "", "2", "A", "---", "1,000", None]
    assert sorted(vals, key=tt.sort_key) == ["2", "10", "1,000", "A", "b", "", "---", None]


# ── Excel ──────────────────────────────────────────────────────────

def test_export_xlsx_layout(tmp_path):
    path = tmp_path / "out.xlsx"
    tt.export_xlsx(path, PROPS, [("G1", ["100", "1時間", "12.5"]), ("G2", ["0012", "", "---"])],
                   changed={("G1", "P-W")})
    wb = load_workbook(path)
    ws = wb[tt.DATA_SHEET]
    assert [c.value for c in ws[1]] == ["GUID", "幅", "耐火", "面積"]
    assert [c.value for c in ws[2]] == ["G1", 100, "1時間", 12.5]   # 数値は数値で出力
    assert [c.value for c in ws[3]] == ["G2", "0012", None, "---"]  # 先頭 0 は文字列のまま
    assert ws["B2"].fill.start_color.rgb.endswith("FFC7CE")         # 未反映の変更は赤塗り
    assert ws["C2"].fill.fill_type is None
    assert ws.freeze_panes == "B2" and ws.auto_filter.ref == "A1:D3"
    assert wb[tt.PROPS_SHEET].sheet_state == "hidden"


def test_export_then_read_roundtrip(tmp_path):
    path = tmp_path / "rt.xlsx"
    rows = [("G1", ["100", "1時間", "12.5"]), ("G2", ["0012", "", "---"])]
    tt.export_xlsx(path, PROPS, rows)
    values, unknown = tt.read_xlsx(path, PROPS)
    assert unknown == []
    assert values == {
        ("G1", "P-W"): "100", ("G1", "P-F"): "1時間", ("G1", "P-RO"): "12.5",
        ("G2", "P-W"): "0012", ("G2", "P-F"): "", ("G2", "P-RO"): "---",
    }


def test_read_uses_hidden_props_sheet_for_duplicate_names(tmp_path):
    props = [{"guid": "A", "name": "名前"}, {"guid": "B", "name": "名前"}]
    path = tmp_path / "dup.xlsx"
    tt.export_xlsx(path, props, [("G1", ["x", "y"])])
    values, _ = tt.read_xlsx(path, props)
    assert values == {("G1", "A"): "x", ("G1", "B"): "y"}


def test_read_plain_sheet_by_header_names(tmp_path):
    """_props シートの無い、ユーザーが手作りした表も見出し名で読める。"""
    wb = Workbook()
    ws = wb.active
    ws.append(["幅", "GUID", "不明な列", None])
    ws.append([150.0, "G1", "?", "x"])
    ws.append([True, "G2"])
    ws.append([1, None])           # GUID 空の行は無視
    path = tmp_path / "plain.xlsx"
    wb.save(path)
    values, unknown = tt.read_xlsx(path, PROPS)
    assert values == {("G1", "P-W"): "150", ("G2", "P-W"): "True"}
    assert unknown == ["不明な列"]


def test_read_requires_guid_column(tmp_path):
    wb = Workbook()
    wb.active.append(["幅"])
    path = tmp_path / "noguid.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="GUID"):
        tt.read_xlsx(path, PROPS)


def test_read_empty_sheet(tmp_path):
    path = tmp_path / "empty.xlsx"
    Workbook().save(path)
    assert tt.read_xlsx(path, PROPS) == ({}, [])


# ── 置換 ───────────────────────────────────────────────────────────

@pytest.mark.parametrize("value, find, repl, kw, expected", [
    ("壁-RC-01", "RC", "S", {}, "壁-S-01"),
    ("rc rc", "RC", "S", {}, "S S"),                                  # 既定は大文字小文字を無視
    ("rc RC", "RC", "S", {"match_case": True}, "rc S"),
    ("abc", "x", "y", {}, None),                                      # 一致なし
    ("RC", "RC", "S", {"whole_cell": True}, "S"),
    ("RC-1", "RC", "S", {"whole_cell": True}, None),                  # 部分一致は対象外
    ("", "", "未設定", {"whole_cell": True}, "未設定"),               # 空欄セルの置換
    ("abc", "", "x", {}, None),                                       # 空検索（部分一致）は何もしない
    ("W-12", r"W-(\d+)", r"壁\1", {"regex": True}, "壁12"),
    ("W-12", r"W-(\d+)", r"壁\1", {"regex": True, "whole_cell": True}, "壁12"),
    ("a.b", ".", "!", {}, "a!b"),                                     # 正規表現オフなら . は文字
    ("a", "a", r"\1", {}, r"\1"),                                     # 正規表現オフなら \ もそのまま
    (None, "", "x", {"whole_cell": True}, "x"),
    (100, "1", "2", {}, "200"),
])
def test_replace_text(value, find, repl, kw, expected):
    assert tt.replace_text(value, find, repl, **kw) == expected


def test_replace_text_invalid_regex():
    with pytest.raises(ValueError, match="正規表現"):
        tt.replace_text("a", "(", "x", regex=True)
